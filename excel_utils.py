"""Excel (XLSX) reading and writing utilities."""

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook


def read_excel(
    path: str | Path,
    sheet_name: str | None = None,
    header_row: int = 1,
) -> list[dict[str, Any]]:
    """Read Excel file into list of dicts.

    Args:
        path: Path to XLSX file
        sheet_name: Sheet to read (default: active sheet)
        header_row: Row containing column headers

    Returns:
        List of dicts, one per row
    """
    wb = load_workbook(path, read_only=True, data_only=True)

    if sheet_name:
        ws = wb[sheet_name]
    else:
        ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(h) if h else f"col_{i}" for i, h in enumerate(rows[header_row - 1])]

    result = []
    for row in rows[header_row:]:
        row_dict = dict(zip(headers, row, strict=False))
        result.append(row_dict)

    wb.close()
    return result


def write_excel(
    path: str | Path,
    data: list[dict[str, Any]],
    sheet_name: str = "Sheet1",
    headers: list[str] | None = None,
) -> None:
    """Write list of dicts to Excel file.

    Args:
        path: Output path
        data: List of dicts to write
        sheet_name: Name of the sheet
        headers: Column headers (auto-detect if None)
    """
    if not data:
        return

    if headers is None:
        headers = list(data[0].keys())

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    # Write headers
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)

    # Write data
    for row_idx, row_data in enumerate(data, 2):
        for col_idx, header in enumerate(headers, 1):
            ws.cell(row=row_idx, column=col_idx, value=row_data.get(header))

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def get_sheet_names(path: str | Path) -> list[str]:
    """Get list of sheet names in Excel file."""
    wb = load_workbook(path, read_only=True)
    names = wb.sheetnames
    wb.close()
    return names


def excel_to_csv(
    excel_path: str | Path,
    csv_path: str | Path,
    sheet_name: str | None = None,
) -> None:
    """Convert Excel sheet to CSV."""
    import csv

    data = read_excel(excel_path, sheet_name=sheet_name)
    if not data:
        return

    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=data[0].keys())
        writer.writeheader()
        writer.writerows(data)
